# ====== ArgoFarm — PPTX Generator Script ======
# This script programmatically builds a professional, high-quality 16:9 PowerPoint presentation
# summarizing the aspects of the ArgoFarm platform, without emojis and in simple language.

import sys
import os
import subprocess

# Auto-install python-pptx if missing
try:
    import pptx
except ImportError:
    print("[INFO] python-pptx is not installed. Attempting to install it via pip...")
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        import pptx
        print("[SUCCESS] python-pptx installed successfully.")
    except Exception as e:
        print(f"[ERROR] Failed to automatically install python-pptx: {e}")
        print("Please run 'pip install python-pptx' manually and re-run this script.")
        sys.exit(1)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Set to widescreen 16:9 aspect ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    c_white = RGBColor(255, 255, 255)
    c_dark_green = RGBColor(21, 128, 61)   # #15803d
    c_primary_green = RGBColor(22, 163, 74) # #16a34a
    c_light_green = RGBColor(240, 253, 244) # #f0fdf4
    c_text_dark = RGBColor(15, 23, 42)      # #0f172a
    c_text_muted = RGBColor(100, 116, 139)  # #64748b
    c_accent_orange = RGBColor(217, 119, 6) # #d97706
    c_light_gray = RGBColor(248, 250, 252)  # #f8fafc

    # Helper: Set slide background
    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper: Add title to slide
    def add_slide_header(slide, title_text, is_dark_bg=False):
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = c_white if is_dark_bg else c_dark_green
        
        # Add decorative line under header
        if not is_dark_bg:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.73), Inches(0.04))
            shape.fill.solid()
            shape.fill.fore_color.rgb = c_primary_green
            shape.line.fill.background() # No border

    # Helper: Add text container
    def add_textbox(slide, left, top, width, height, text_lines, font_size=15, is_bullet=True):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        
        for idx, line in enumerate(text_lines):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            
            # Formatting logic for pseudo-bold tags "**"
            remaining_text = line
            p.font.name = "Arial"
            p.font.size = Pt(font_size)
            
            if is_bullet:
                p.level = 0
                
            # Parse bold parts
            while "**" in remaining_text:
                start_bold = remaining_text.find("**")
                end_bold = remaining_text.find("**", start_bold + 2)
                if end_bold == -1:
                    break
                
                # Add normal text before bold
                if start_bold > 0:
                    run = p.add_run()
                    run.text = remaining_text[:start_bold]
                
                # Add bold text
                run_bold = p.add_run()
                run_bold.text = remaining_text[start_bold+2:end_bold]
                run_bold.font.bold = True
                run_bold.font.color.rgb = c_text_dark
                
                remaining_text = remaining_text[end_bold+2:]
            
            if remaining_text:
                run = p.add_run()
                run.text = remaining_text
                
        return box

    # ====== SLIDE 1: Cover (Dark Green Background) ======
    slide_layout = prs.slide_layouts[6] # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_dark_green)
    
    # Outer decorative card border
    border_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.33), Inches(6.5))
    border_shape.fill.background()
    border_shape.line.color.rgb = RGBColor(134, 239, 172) # Light green border
    border_shape.line.width = Pt(3)
    
    logo_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.33), Inches(1.0))
    p = logo_box.text_frame.paragraphs[0]
    p.text = "ARGOFARM"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(187, 247, 208)
    
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.33), Inches(1.2))
    p = title_box.text_frame.paragraphs[0]
    p.text = "ArgoFarm"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = c_white
    
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(11.33), Inches(1.0))
    p = sub_box.text_frame.paragraphs[0]
    p.text = "Agricultural Advisory & Crop Diagnostics Platform for Pakistan"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(187, 247, 208) # Mint green
    
    desc_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.33), Inches(0.6))
    p = desc_box.text_frame.paragraphs[0]
    p.text = "Project Overview & System Architecture"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.color.rgb = c_white
    
    meta_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.33), Inches(1.0))
    tf = meta_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Platform: Single-Page Web App (SPA)   |   Target Region: Pakistan   |   Languages: English & Urdu"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(187, 247, 208)

    # ====== SLIDE 2: Project Overview ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "1. Project Overview & Impact")
    
    # Left Card: Problem Statement
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(1.0), Inches(1.8), Inches(5.2), Inches(0.6), ["**The Challenge:**"], font_size=20, is_bullet=False)
    add_textbox(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Crop Yields:** Many farmers face reduced crop yields because plant diseases are not identified early enough.",
        "**Advisor Access:** Expert agricultural support is often limited or hard to reach in rural farming villages.",
        "**Weather Changes:** Climate variations impact planting schedules without timely local advice.",
        "**Market Pricing:** Lack of transparent crop prices makes it difficult to negotiate fair bulk sales."
    ], font_size=14, is_bullet=True)
    
    # Right Card: Solution Goals
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_light_green
    card2.line.color.rgb = RGBColor(187, 247, 208)
    
    add_textbox(slide, Inches(7.1), Inches(1.8), Inches(5.2), Inches(0.6), ["**The Solution:**"], font_size=20, is_bullet=False)
    add_textbox(slide, Inches(7.1), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Crop Diagnostics:** Upload leaf photos to identify plant health issues quickly.",
        "**Bilingual Support:** Ask questions naturally in English or Urdu via text or voice.",
        "**Regional Health Trends:** Display local crop warnings and advisories on an interactive map.",
        "**Farming Planners:** Generate simple 4-month crop timelines and wholesale pricing suggestions."
    ], font_size=14, is_bullet=True)

    # ====== SLIDE 3: Tech Stack ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "2. Technology Stack")
    
    # 4 Cards for different layers
    layers = [
        {"title": "Backend API Server", "tech": ["Python / Flask Server", "mysql-connector API helper", "PyJWT & Bcrypt logins", "scikit-learn crop recommender", "Pillow image checker"]},
        {"title": "Frontend Interface", "tech": ["Vanilla JS Single-Page App", "Vite build tool", "Leaflet.js map display", "Chart.js history logs", "CSS Variables styling"]},
        {"title": "AI Integrations", "tech": ["Google Gemini", "Groq LLaMA", "Groq Whisper", "Open-Meteo API"]},
        {"title": "Deployment", "tech": ["Docker Compose services", "Nginx gateway proxy", "Netlify host routing", "VPS server environment", "Git version manager"]}
    ]
    
    for idx, layer in enumerate(layers):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.6 + row * 2.7)
        
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = c_white
        card.line.color.rgb = RGBColor(226, 232, 240)
        
        add_textbox(slide, left + Inches(0.2), top + Inches(0.2), Inches(5.2), Inches(0.4), [f"**{layer['title']}**"], font_size=16, is_bullet=False)
        add_textbox(slide, left + Inches(0.2), top + Inches(0.6), Inches(5.2), Inches(1.6), layer["tech"], font_size=13, is_bullet=True)

    # ====== SLIDE 4: System Architecture ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "3. How the Platform Works")
    
    # Descriptive list
    add_textbox(slide, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.0), [
        "ArgoFarm uses a simple client-server layout designed for easy access, account security, and offline support."
    ], font_size=16, is_bullet=False)
    
    arch_flow = [
        "**1. User Interface (Browser):** A web-based interface optimized for mobile phone screens in the field.",
        "**2. Web Proxy Gateway:** Safely routes requests from the user interface to the core server.",
        "**3. Application Server:** Manages the core database, security checks, and platform operations.",
        "**4. Database Layer:** Securely stores user profiles, crop planning data, and chat logs.",
        "**5. AI Services Layer:** Coordinates voice transcripts, leaf scans, and market price evaluations."
    ]
    add_textbox(slide, Inches(0.8), Inches(2.6), Inches(11.7), Inches(4.5), arch_flow, font_size=14, is_bullet=True)

    # ====== SLIDE 5: Core Features (Scanning & Chat) ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "4. Disease Diagnostics & Advisor")
    
    # Scanning Details
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(1.0), Inches(1.8), Inches(5.2), Inches(0.6), ["**Crop Disease Scanning**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Photo Scanning:** Leaf images are analyzed by agricultural vision models to detect health issues.",
        "**Quality Guard:** Blurry or non-crop photos are flagged to ensure accurate readings.",
        "**Simple Remedies:** Get practical organic treatments, pesticide controls, and preventive steps.",
        "**System Fallbacks:** Returns standard crop health advice if the cloud service is offline."
    ], font_size=13, is_bullet=True)
    
    # Conversational Chat Details
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_white
    card2.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(7.1), Inches(1.8), Inches(5.2), Inches(0.6), ["**CropMind Chat Advisor**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(7.1), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Farming Assistant:** Converse with a digital agronomist helper that focuses strictly on agricultural topics.",
        "**Voice Translation:** Speak questions aloud in local languages for easy hands-free use.",
        "**Auto Forms:** Mentions of crop planning or soil queries automatically open the correct input form.",
        "**Context Memory:** Remembers recent exchanges in the conversation to support continuous discussions."
    ], font_size=13, is_bullet=True)

    # ====== SLIDE 6: Soil Recommendation, Planning, Marketplace ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "5. Soil Suggestions, Planner & Market")
    
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    add_textbox(slide, Inches(0.9), Inches(1.8), Inches(3.5), Inches(0.5), ["**Soil Recommendations**"], font_size=16, is_bullet=False)
    add_textbox(slide, Inches(0.9), Inches(2.4), Inches(3.5), Inches(4.0), [
        "**Soil Inputs:** Analyzes nitrogen, potassium, pH, and rainfall to select the best crop.",
        "**Local Model:** Ingests soil metrics locally to recommend ideal crop choices.",
        "**Custom Guides:** Generates a basic layout for fertilizer and watering schedules."
    ], font_size=13, is_bullet=True)
    
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_white
    card2.line.color.rgb = RGBColor(226, 232, 240)
    add_textbox(slide, Inches(4.9), Inches(1.8), Inches(3.5), Inches(0.5), ["**AI Crop Planner**"], font_size=16, is_bullet=False)
    add_textbox(slide, Inches(4.9), Inches(2.4), Inches(3.5), Inches(4.0), [
        "**Inputs:** Ingests crop type, acreage, water availability, and budget.",
        "**4-Month Timeline:** Creates a detailed timeline from soil prep to final harvest.",
        "**English & Urdu:** Displays guides in the farmer's preferred language."
    ], font_size=13, is_bullet=True)

    card3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.0))
    card3.fill.solid()
    card3.fill.fore_color.rgb = c_light_green
    card3.line.color.rgb = RGBColor(187, 247, 208)
    add_textbox(slide, Inches(8.9), Inches(1.8), Inches(3.5), Inches(0.5), ["**Wholesale Marketplace**"], font_size=16, is_bullet=False)
    add_textbox(slide, Inches(8.9), Inches(2.4), Inches(3.5), Inches(4.0), [
        "**Crop Listings:** Post wholesale listings with custom pricing and location.",
        "**Deal Analyzer:** Calculates estimated profit margins and suggests negotiation tips.",
        "**Price Suggestions:** Compares inputs with local mandi averages to recommend fair prices."
    ], font_size=13, is_bullet=True)

    # ====== SLIDE 7: Database Design ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "6. Database and Storage")
    
    add_textbox(slide, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.8), [
        "Relational database layout with MySQL 8.0 for production and SQLite for development. Simple table structures hold key data:"
    ], font_size=15, is_bullet=False)
    
    # Table details
    table_details = [
        "**1. users:** Handles account credentials, region settings, and daily AI limits.",
        "**2. scans:** Saves diagnostic logs, leaf photo links, and treatment guides.",
        "**3. chat_history:** Keeps user conversation histories for the digital assistant.",
        "**4. community_posts & comments:** Manages discussion boards and replies.",
        "**5. post_likes:** Tracks forum likes to ensure fair discussion voting.",
        "**6. marketplace_items:** Keeps wholesale crop listings posted by farmers."
    ]
    add_textbox(slide, Inches(0.8), Inches(2.5), Inches(11.7), Inches(4.5), table_details, font_size=13, is_bullet=True)

    # ====== SLIDE 8: API Rate Limiting ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "7. Rate Limiting & Stability")
    
    # Column 1: Architecture
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(1.0), Inches(1.8), Inches(5.2), Inches(0.5), ["**Rate Limiting Controls**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Security Wrappers:** Uses decorator functions to protect backend API routes.",
        "**Memory Store:** Keeps temporary counts of route access by time windows.",
        "**Client Check:** Identifies users by account token or device IP address.",
        "**Short & Long Term:** Enforces both per-minute and daily access restrictions.",
        "**Route Cooldown:** Temporarily suspends access if limits are exceeded, providing a wait timer."
    ], font_size=13, is_bullet=True)
    
    # Column 2: Specific Limits
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_white
    card2.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(7.1), Inches(1.8), Inches(5.2), Inches(0.5), ["**Default Usage Quotas**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(7.1), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Leaf Scan:** 3 requests per minute, 10 requests per day.",
        "**Chat Advisor:** 10 requests per minute, 100 requests per day.",
        "**Voice Translation:** 5 requests per minute, 50 requests per day.",
        "**Weather Alerts:** 5 requests per minute, 30 requests per day.",
        "**Price Suggestions:** 5 requests per minute, 30 requests per day.",
        "**Planner & AI Advice:** 5 requests per minute, 30 requests per day."
    ], font_size=13, is_bullet=True)

    # ====== SLIDE 9: Admin Management Console ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "8. Admin Controls & Governance")
    
    # Left Card: User control
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(1.0), Inches(1.8), Inches(5.2), Inches(0.5), ["**Admin Panel Features**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**User Management:** View user accounts, email registration, and regional settings.",
        "**Admin Privilege:** Safely grant or demote administrator status for users.",
        "**Auto Bootstrap:** Promotes the first user in the database to admin automatically.",
        "**User Deletion:** Permanently delete inactive accounts and related histories."
    ], font_size=13, is_bullet=True)
    
    # Right Card: Custom AI Overrides
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_white
    card2.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(7.1), Inches(1.8), Inches(5.2), Inches(0.5), ["**Custom Quota Overrides**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(7.1), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**AI Limit Customization:** Set custom daily quotas for individual accounts.",
        "**Access Disabling:** Setting daily limits to zero disables all AI tool options.",
        "**Access Warnings:** Disabled accounts receive clear warnings on blocked pages.",
        "**Self Lockout Safeguard:** Admin accounts are protected from self-demotion or self-deletion."
    ], font_size=13, is_bullet=True)

    # ====== SLIDE 10: Platform Security ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "9. Security Measures")
    
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(1.0), Inches(1.8), Inches(5.2), Inches(0.5), ["**Account Security**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Password Hashing:** Hashes user passwords securely using Bcrypt.",
        "**Token Logins:** Signed Web Tokens manage secure user sessions.",
        "**Auto Logout:** Clears token keys if sessions expire.",
        "**CORS Settings:** Limits server entry to registered frontend origins."
    ], font_size=13, is_bullet=True)
    
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_white
    card2.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(7.1), Inches(1.8), Inches(5.2), Inches(0.5), ["**Data Security**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(7.1), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**SQL Protection:** Uses parameter bindings to prevent database injection.",
        "**Email Sanitization:** Enforces lowercase checking to prevent account duplicates.",
        "**Key Isolation:** External API credentials are kept in secure local environments.",
        "**Forum Likes Mapping:** Uses mapping parameters to restrict users to one vote per post."
    ], font_size=13, is_bullet=True)

    # ====== SLIDE 11: Deployment & DevOps ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_light_gray)
    add_slide_header(slide, "10. Deployment & Hosting")
    
    # Left Card: Docker Compose
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_white
    card1.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(1.0), Inches(1.8), Inches(5.2), Inches(0.5), ["**Containerized Stack**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Database container:** MySQL production container with health checks and data volumes.",
        "**Backend container:** Flask application running core routes.",
        "**Proxy router:** Nginx gateway managing SSL and request directions."
    ], font_size=13, is_bullet=True)
    
    # Right Card: CDN & Proxy Rules
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_white
    card2.line.color.rgb = RGBColor(226, 232, 240)
    
    add_textbox(slide, Inches(7.1), Inches(1.8), Inches(5.2), Inches(0.5), ["**Cloud Hosting**"], font_size=18, is_bullet=False)
    add_textbox(slide, Inches(7.1), Inches(2.4), Inches(5.2), Inches(4.0), [
        "**Frontend CDN:** Hosts frontend file resources on a fast content network.",
        "**Proxy Settings:** Redirects API requests smoothly to bypass security errors."
    ], font_size=13, is_bullet=True)

    # ====== SLIDE 12: Conclusion (Dark Green Background) ======
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, c_dark_green)
    
    border_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.33), Inches(6.5))
    border_shape.fill.background()
    border_shape.line.color.rgb = RGBColor(134, 239, 172)
    border_shape.line.width = Pt(3)
    
    logo_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.33), Inches(1.0))
    p = logo_box.text_frame.paragraphs[0]
    p.text = "ARGOFARM"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(187, 247, 208)
    
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(11.33), Inches(1.2))
    p = title_box.text_frame.paragraphs[0]
    p.text = "Thank You"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = c_white
    
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.33), Inches(1.0))
    p = sub_box.text_frame.paragraphs[0]
    p.text = "ArgoFarm — Agricultural Advisory & Crop Diagnostics Platform"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(187, 247, 208)
    
    desc_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.8), Inches(11.33), Inches(0.6))
    p = desc_box.text_frame.paragraphs[0]
    p.text = "Supporting Pakistani farmers through simple and accessible technology."
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.color.rgb = c_white

    # Save presentation with lock-prevention fallback
    output_filename = "ArgoFarm_Presentation.pptx"
    try:
        prs.save(output_filename)
        print(f"[SUCCESS] PowerPoint presentation saved as '{output_filename}'.")
    except PermissionError:
        fallback_filename = "ArgoFarm_Presentation_New.pptx"
        print(f"[WARNING] '{output_filename}' is currently locked (likely open in PowerPoint). Saving copy as '{fallback_filename}'...")
        prs.save(fallback_filename)
        print(f"[SUCCESS] PowerPoint presentation saved as '{fallback_filename}'.")

if __name__ == "__main__":
    create_presentation()
